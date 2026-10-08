#!/usr/bin/env python3
"""
I testi critici d'artista, pronti da stampare e da condividere.

Legge la scheda di un artista, testi-critici/artisti/<nome-cognome>/scheda.json,
e produce nella sua cartella:

- presentazione — la sintesi critica in A5, da consegnare all'artista e da
  stampare per la sala
- una didascalia in A6 per ogni opera analizzata, da mettere accanto al quadro

ognuna in tre forme: pdf/ per la stampa, png/ per mandarla in chat o sui
social, word/ per ritoccarla a mano prima di stampare.

    py strumenti\\critica.py nome-cognome      un artista
    py strumenti\\critica.py                   tutti

La scheda la scrive Claude: qui non c'è niente da comporre a mano.
"""

import json
import re
import sys
import unicodedata
from pathlib import Path

from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph

import pypdfium2
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

CARTELLA = Path(__file__).resolve().parent
RADICE = CARTELLA.parent
ARTISTI = RADICE / "testi-critici" / "artisti"
FONT = CARTELLA / "marchio" / "font"
LOGO = CARTELLA / "marchio" / "logo-nero.png"   # il marchio chiaro vuole il suo fondo nero

A5 = (148 * mm, 210 * mm)
A6_ORIZZONTALE = (148 * mm, 105 * mm)

CARTA = HexColor("#FAF7F0")
INCHIOSTRO = HexColor("#1E1B18")
GRIGIO = HexColor("#6B645B")
ORO = HexColor("#C9A227")   # l'oro antico del logo, lo stesso dei social

# Il PNG ha la risoluzione della stampa: regge anche se lo si porta in tipografia
PNG_DPI = 300

# Word usa i caratteri installati sul computer, non quelli del repository: il
# Garamond arriva insieme a Office, è il parente più stretto del Cormorant e
# così il documento si apre uguale su qualunque PC, senza installare niente.
# PDF e PNG restano in Cormorant: lì il carattere viaggia dentro il file
FONT_WORD = "Garamond"

# Le soglie della procedura: oltre, il testo non sta più nel suo formato
LIMITI = {"presentazione": (500, 600), "didascalia": (150, 300)}

PIEDE = "BE Art Gallery & Creative Lab  ·  Pistoia  ·  beartgallery.eu"

# Un carattere solo, quello delle locandine, nei suoi pesi
for nome, file in {
    "Cormorant": "CormorantGaramond-Regular.ttf",
    "Cormorant-Forte": "CormorantGaramond-SemiBold.ttf",
    "Cormorant-Corsivo": "CormorantGaramond-Italic.ttf",
}.items():
    pdfmetrics.registerFont(TTFont(nome, str(FONT / file)))


def stile(nome, **opzioni) -> ParagraphStyle:
    base = dict(fontName="Cormorant", fontSize=11.5, leading=15.5,
                textColor=INCHIOSTRO, alignment=TA_LEFT)
    base.update(opzioni)
    return ParagraphStyle(nome, **base)


def nome_file(testo: str) -> str:
    """«Il colore come linguaggio» → il-colore-come-linguaggio"""
    piano = unicodedata.normalize("NFKD", testo).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", piano.lower()).strip("-") or "opera"


def sottotitolo(scheda: dict) -> str:
    """La riga sotto il nome: dove vive l'artista, o — per una serie senza
    autore, una collezione — il "sottotitolo" scritto nella scheda."""
    if scheda.get("sottotitolo"):
        return scheda["sottotitolo"]
    if scheda.get("vive_e_lavora"):
        return f"vive e lavora a {scheda['vive_e_lavora']}"
    return ""


def nome_opera(opera: dict) -> str:
    """Il nome dei file di un'opera. Due opere con lo stesso titolo — «Paesaggio»,
    «Senza titolo» — si distinguono con "file" nella scheda, o si sovrascrivono."""
    return opera.get("file") or nome_file(opera["titolo"])


def controlla(etichetta: str, testo: str, tipo: str) -> None:
    minimo, massimo = LIMITI[tipo]
    n = len(testo)
    if not minimo <= n <= massimo:
        print(f"   ⚠ {etichetta}: {n} caratteri, la procedura ne vuole {minimo}–{massimo}")


def misura(testo: str, st: ParagraphStyle, larghezza: float) -> float:
    """L'altezza che un paragrafo occuperebbe, senza disegnarlo."""
    p = Paragraph(testo.replace("'", "\u2019").replace("\n", "<br/>"), st)
    return p.wrap(larghezza, 10_000)[1]


def scrivi(c: Canvas, testo: str, st: ParagraphStyle, x: float, y: float, larghezza: float) -> float:
    """Posa un paragrafo con il bordo superiore in y; restituisce dove finisce."""
    # Apostrofi tipografici: quello dritto della tastiera sulla carta si nota
    testo = testo.replace("'", "\u2019").replace("\n", "<br/>")
    p = Paragraph(testo, st)
    _, altezza = p.wrap(larghezza, 10_000)
    p.drawOn(c, x, y - altezza)
    return y - altezza


def fondo(c: Canvas, formato) -> None:
    c.setFillColor(CARTA)
    c.rect(0, 0, *formato, stroke=0, fill=1)


def filetto(c: Canvas, x: float, y: float, lunghezza: float, spessore=0.6) -> None:
    c.setStrokeColor(ORO)
    c.setLineWidth(spessore)
    c.line(x, y, x + lunghezza, y)


def logo(c: Canvas, x: float, y_alto: float, larghezza: float) -> None:
    altezza = larghezza * 616 / 888
    c.drawImage(str(LOGO), x, y_alto - altezza, larghezza, altezza, mask="auto")


def piede(c: Canvas, larghezza_pagina: float, margine: float, base: float, corpo=7.5) -> None:
    filetto(c, margine, base + 9, larghezza_pagina - 2 * margine, 0.4)
    c.setFont("Cormorant", corpo)
    c.setFillColor(GRIGIO)
    c.drawCentredString(larghezza_pagina / 2, base, PIEDE)


def presentazione(scheda: dict, uscita: Path) -> Path:
    testi = scheda["presentazione"]
    controlla("presentazione (it)", testi["it"], "presentazione")

    larghezza, altezza = A5
    m = 16 * mm
    utile = larghezza - 2 * m
    file = uscita / "presentazione.pdf"
    c = Canvas(str(file), pagesize=A5)
    c.setTitle(f"{scheda['nome']} — testo critico")
    c.setAuthor("BE Art Gallery & Creative Lab")
    fondo(c, A5)

    logo(c, m, altezza - m, 24 * mm)
    y = altezza - m - 30 * mm

    y = scrivi(c, scheda["nome"], stile("nome", fontName="Cormorant-Forte",
               fontSize=25, leading=28), m, y, utile)
    if sottotitolo(scheda):
        y = scrivi(c, sottotitolo(scheda),
                   stile("luogo", fontName="Cormorant-Corsivo", fontSize=11,
                         leading=14, textColor=GRIGIO), m, y - 2, utile)
    y -= 5 * mm
    filetto(c, m, y, 18 * mm, 1)
    y -= 6 * mm

    # I testi di una serie possono superare i 600 caratteri: invece di
    # sbordare sul piede, il corpo scende di misura quanto basta
    fondo_utile = m + 9 + 5 * mm
    for k in (1, 0.95, 0.9, 0.85, 0.8, 0.75):
        corpo = stile("it", fontSize=12 * k, leading=16.5 * k, alignment=TA_JUSTIFY)
        inglese = stile("en", fontName="Cormorant-Corsivo", fontSize=10.5 * k,
                        leading=14 * k, textColor=GRIGIO, alignment=TA_JUSTIFY)
        serve = misura(testi["it"], corpo, utile)
        if testi.get("en"):
            serve += 11 * mm * k + misura(testi["en"], inglese, utile)
        if y - serve >= fondo_utile:
            break

    y = scrivi(c, testi["it"], corpo, m, y, utile)

    if testi.get("en"):
        y -= 6 * mm * k
        filetto(c, m, y, 8 * mm, 0.6)
        y -= 5 * mm * k
        scrivi(c, testi["en"], inglese, m, y, utile)

    piede(c, larghezza, m, m)
    c.save()
    return file


def didascalia(scheda: dict, opera: dict, uscita: Path) -> Path:
    testi = opera["didascalia"]
    controlla(f"didascalia «{opera['titolo']}» (it)", testi["it"], "didascalia")

    larghezza, altezza = A6_ORIZZONTALE
    m = 11 * mm
    utile = larghezza - 2 * m
    file = uscita / f"opera-{nome_opera(opera)}.pdf"
    c = Canvas(str(file), pagesize=A6_ORIZZONTALE)
    c.setTitle(f"{scheda['nome']} — {opera['titolo']}")
    c.setAuthor("BE Art Gallery & Creative Lab")
    fondo(c, A6_ORIZZONTALE)

    # Il cartellino da mostra: chi, cosa, come — poi la lettura
    logo(c, larghezza - m - 15 * mm, altezza - m + 2 * mm, 15 * mm)
    y = altezza - m
    y = scrivi(c, scheda["nome"], stile("artista", fontName="Cormorant-Forte",
               fontSize=11, leading=13), m, y, utile - 18 * mm)
    y = scrivi(c, opera["titolo"], stile("titolo", fontName="Cormorant-Corsivo",
               fontSize=16, leading=19), m, y - 1, utile - 18 * mm)
    dati = "  ·  ".join(str(opera[k]) for k in ("anno", "tecnica", "misure", "provenienza") if opera.get(k))
    if dati:
        y = scrivi(c, dati, stile("dati", fontSize=9, leading=11.5,
                   textColor=GRIGIO), m, y - 1, utile)
    y -= 3.5 * mm
    filetto(c, m, y, 12 * mm, 0.8)
    y -= 3.5 * mm

    y = scrivi(c, testi["it"], stile("dit", fontSize=10.5, leading=13.5,
               alignment=TA_JUSTIFY), m, y, utile)
    if testi.get("en"):
        scrivi(c, testi["en"], stile("den", fontName="Cormorant-Corsivo",
               fontSize=9, leading=11.5, textColor=GRIGIO,
               alignment=TA_JUSTIFY), m, y - 2.5 * mm, utile)

    piede(c, larghezza, m, 7 * mm, corpo=6.5)
    c.save()
    return file


def png(pdf: Path, uscita: Path) -> Path:
    """La stessa pagina del PDF, in immagine: niente da impaginare due volte."""
    file = uscita / f"{pdf.stem}.png"
    documento = pypdfium2.PdfDocument(str(pdf))
    documento[0].render(scale=PNG_DPI / 72).to_pil().save(file, dpi=(PNG_DPI, PNG_DPI))
    documento.close()
    return file


# ----------------------------------------------------------------------------
# Word: la stessa pagina in forma modificabile
# ----------------------------------------------------------------------------

def _rgb(colore) -> RGBColor:
    return RGBColor.from_string(colore.hexval()[2:].upper())


def _paragrafo(doc, testo="", corpo=11.5, colore=INCHIOSTRO, grassetto=False,
               corsivo=False, giustificato=False, centrato=False, dopo=0, interlinea=None):
    p = doc.add_paragraph()
    if giustificato:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if centrato:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(dopo)
    if interlinea:
        p.paragraph_format.line_spacing = Pt(interlinea)
    if testo:
        r = p.add_run(testo.replace("'", "\u2019"))
        r.font.name = FONT_WORD
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_WORD)
        r.font.size = Pt(corpo)
        r.font.color.rgb = _rgb(colore)
        r.bold = grassetto
        r.italic = corsivo
    return p


def _filetto(paragrafo, colore=ORO, spessore=6) -> None:
    """Il filetto oro come bordo inferiore del paragrafo (spessore in ottavi di punto)."""
    bordi = OxmlElement("w:pBdr")
    sotto = OxmlElement("w:bottom")
    for chiave, valore in {"w:val": "single", "w:sz": str(spessore),
                           "w:space": "4", "w:color": colore.hexval()[2:].upper()}.items():
        sotto.set(qn(chiave), valore)
    bordi.append(sotto)
    paragrafo._p.get_or_add_pPr().append(bordi)


def _documento(formato, margine, titolo):
    doc = Document()
    doc.core_properties.title = titolo
    doc.core_properties.author = "BE Art Gallery & Creative Lab"
    sezione = doc.sections[0]
    sezione.page_width, sezione.page_height = Mm(formato[0] / mm), Mm(formato[1] / mm)
    for lato in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sezione, lato, Mm(margine))
    sezione.footer_distance = Mm(margine * 0.6)
    piede = sezione.footer.paragraphs[0]
    piede.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = piede.add_run(PIEDE)
    r.font.name, r.font.size, r.font.color.rgb = FONT_WORD, Pt(7.5), _rgb(GRIGIO)
    return doc


def word_presentazione(scheda: dict, uscita: Path) -> Path:
    testi = scheda["presentazione"]
    doc = _documento(A5, 16, f"{scheda['nome']} — testo critico")
    doc.add_paragraph().add_run().add_picture(str(LOGO), width=Mm(24))
    _paragrafo(doc, dopo=18)
    _paragrafo(doc, scheda["nome"], corpo=25, grassetto=True)
    luogo = _paragrafo(doc, sottotitolo(scheda),
                       corpo=11, colore=GRIGIO, corsivo=True, dopo=14)
    _filetto(luogo)
    _paragrafo(doc, testi["it"], corpo=12, giustificato=True, interlinea=16.5, dopo=18)
    if testi.get("en"):
        _paragrafo(doc, testi["en"], corpo=10.5, colore=GRIGIO, corsivo=True,
                   giustificato=True, interlinea=14)
    file = uscita / "presentazione.docx"
    doc.save(file)
    return file


def word_didascalia(scheda: dict, opera: dict, uscita: Path) -> Path:
    testi = opera["didascalia"]
    doc = _documento(A6_ORIZZONTALE, 11, f"{scheda['nome']} — {opera['titolo']}")
    _paragrafo(doc, scheda["nome"], corpo=11, grassetto=True)
    _paragrafo(doc, opera["titolo"], corpo=16, corsivo=True)
    dati = "  ·  ".join(str(opera[k]) for k in ("anno", "tecnica", "misure", "provenienza") if opera.get(k))
    _filetto(_paragrafo(doc, dati, corpo=9, colore=GRIGIO, dopo=8))
    _paragrafo(doc, testi["it"], corpo=10.5, giustificato=True, interlinea=13.5, dopo=6)
    if testi.get("en"):
        _paragrafo(doc, testi["en"], corpo=9, colore=GRIGIO, corsivo=True,
                   giustificato=True, interlinea=11.5)
    file = uscita / f"opera-{nome_opera(opera)}.docx"
    doc.save(file)
    return file


def artista(cartella: Path) -> list[Path]:
    scheda = json.loads((cartella / "scheda.json").read_text(encoding="utf-8"))
    cartelle = {formato: cartella / formato for formato in ("pdf", "png", "word")}
    for c in cartelle.values():
        c.mkdir(exist_ok=True)
    print(f"\n{scheda['nome']}")

    lavori = []
    if scheda.get("presentazione", {}).get("it"):
        lavori.append((presentazione, word_presentazione, (scheda,)))
    for opera in scheda.get("opere", []):
        if opera.get("didascalia", {}).get("it"):
            lavori.append((didascalia, word_didascalia, (scheda, opera)))

    nomi = [nome_opera(o) for o in scheda.get("opere", [])]
    doppi = {n for n in nomi if nomi.count(n) > 1}
    if doppi:
        print(f"   ⚠ più opere finirebbero nello stesso file: {', '.join(sorted(doppi))}"
              " — aggiungi \"file\" nella scheda per distinguerle")
        raise SystemExit(1)

    prodotti = []
    for fai_pdf, fai_word, argomenti in lavori:
        pdf = fai_pdf(*argomenti, cartelle["pdf"])
        prodotti += [pdf, png(pdf, cartelle["png"]), fai_word(*argomenti, cartelle["word"])]
    for file in prodotti:
        print(f"   {file.relative_to(RADICE)}")
    return prodotti


def main() -> None:
    if len(sys.argv) > 1:
        cartelle = [ARTISTI / sys.argv[1]]
        if not (cartelle[0] / "scheda.json").exists():
            print(f"\nNon trovo testi-critici/artisti/{sys.argv[1]}/scheda.json\n")
            raise SystemExit(1)
    else:
        # Le cartelle che iniziano con _ sono modelli, non artisti
        cartelle = sorted(p.parent for p in ARTISTI.glob("*/scheda.json")
                          if not p.parent.name.startswith("_"))
    for cartella in cartelle:
        artista(cartella)
    print()


if __name__ == "__main__":
    main()
